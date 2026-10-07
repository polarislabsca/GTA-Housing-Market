"""Add one month of TRREB Market Watch data to the dashboard JSON and the workbook.

Usage: python scripts/add_month.py 2026 10

Extracts every property type from TRREB/mwYYMM.pdf, checks that the area totals add up,
then merges the month into public/data/market-data.json and appends the detached and
price-range rows to the workbook (renamed to the new latest month). Nothing is written
if an extraction check fails, unless --force is passed. Notes about inconsistencies in
TRREB's own report are printed but don't block the update.
"""
import json
import re
import shutil
import sys

import openpyxl

from generate_dashboard_data import OUTPUT, write_payload
from trreb_extract import PROJECT, check, extract_month, unique_rows

DATA_DIR = OUTPUT.parent
SOURCE_URL = "https://trreb.ca/wp-content/files/market-stats/market-watch/mw{yymm}.pdf"


def merge_json(records, year, month):
    date = f"{year}-{month:02d}-01"
    existing = [r for r in json.loads(OUTPUT.read_text())["records"] if r["date"] != date]
    added = list(unique_rows(records).values())
    payload = write_payload(existing + added)
    return len(added), payload["metadata"]["updatedThrough"]


def extend_table(sheet, table_name, last_row):
    table = sheet.tables[table_name]
    start, end = table.ref.split(":")
    end = re.sub(r"\d+$", str(last_row), end)
    table.ref = f"{start}:{end}"


def detached_sheet_rows(records, year, month):
    url = SOURCE_URL.format(yymm=f"{year % 100:02d}{month:02d}")
    rows = []
    for scope in ("ALL TRREB", "City of Toronto"):
        for r in sorted((r for r in records if r["propertyType"] == "Detached" and r["scope"] == scope),
                        key=lambda r: r["city"]):
            rows.append([year, month, r["city"], scope, r["sales"], r["newListings"], r["activeListings"],
                         r["monthsOfInventory"], r["dollarVolume"], r["averagePrice"], r["medianPrice"],
                         r["saleToList"], r["daysOnMarket"], url])
    return rows


def update_workbook(data, year, month):
    current = sorted(DATA_DIR.glob("TRREB_Detached_Dataset_through_*.xlsx"))[-1]
    target = DATA_DIR / f"TRREB_Detached_Dataset_through_{year}-{month:02d}.xlsx"
    workbook = openpyxl.load_workbook(current)
    charts_before = sum(len(ws._charts) for ws in workbook.worksheets)
    url = SOURCE_URL.format(yymm=f"{year % 100:02d}{month:02d}")

    stats = workbook["Detached_Stats"]
    if any((r[0], r[1]) == (year, month) for r in stats.iter_rows(min_row=2, max_col=2, values_only=True)):
        raise RuntimeError(f"Workbook already has {year}-{month:02d} rows; not appending twice.")
    for row in detached_sheet_rows(data["records"], year, month):
        stats.append(row)
    extend_table(stats, "DetachedStatsTable", stats.max_row)

    ranges = workbook["Price_Range_Detached"]
    for label, count in data["priceRanges"]:
        ranges.append([year, month, label, count, url])
    extend_table(ranges, "PriceRangeDetachedTable", ranges.max_row)

    workbook.save(target)
    charts_after = sum(len(ws._charts) for ws in openpyxl.load_workbook(target).worksheets)
    if target != current:
        current.unlink()
    shutil.copyfile(target, PROJECT / "TRREB_Detached_Dataset.xlsx")
    return target, charts_before, charts_after


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    year, month = int(args[0]), int(args[1])
    data = extract_month(year, month)
    problems = check(data)
    print(f"Extracted {len(data['records'])} rows and {len(data['priceRanges'])} price ranges for {year}-{month:02d}")
    for kind, message in problems:
        print(f"  [{kind}] {message}")
    if any(kind == "extraction" for kind, _ in problems) and "--force" not in sys.argv:
        sys.exit("Nothing written. Fix the extraction problems above or rerun with --force.")
    added, latest = merge_json(data["records"], year, month)
    print(f"market-data.json: {added} records for {year}-{month:02d}; data now runs through {latest[:7]}")
    target, before, after = update_workbook(data, year, month)
    print(f"Workbook: {target.name} (charts before/after save: {before}/{after}), copied to top-level TRREB_Detached_Dataset.xlsx")


if __name__ == "__main__":
    main()
