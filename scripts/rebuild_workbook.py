"""Rewrite the workbook's Detached_Stats and Price_Range_Detached sheets from every report.

Usage: python scripts/rebuild_workbook.py

Use after fixing the extractor; add_month.py appends one month at a time. Analysis_2026
is left as it is: its formulas read the two tables, so they pick up the new values.
"""
import shutil

import openpyxl

from add_month import DATA_DIR, SOURCE_URL, detached_sheet_rows, extend_table
from generate_dashboard_data import report_months
from trreb_extract import PROJECT, extract_month


def replace_rows(sheet, table_name, rows):
    old_rows = {tuple(row) for row in sheet.iter_rows(min_row=2, values_only=True)}
    changed = sum(1 for row in rows if tuple(row) not in old_rows)
    sheet.delete_rows(2, sheet.max_row)
    for row in rows:
        sheet.append(row)
    extend_table(sheet, table_name, len(rows) + 1)
    return changed


def main():
    path = sorted(DATA_DIR.glob("TRREB_Detached_Dataset_through_*.xlsx"))[-1]
    workbook = openpyxl.load_workbook(path)
    charts_before = sum(len(ws._charts) for ws in workbook.worksheets)
    detached, ranges = [], []
    for year, month in report_months():
        data = extract_month(year, month)
        detached += detached_sheet_rows(data["records"], year, month)
        url = SOURCE_URL.format(yymm=f"{year % 100:02d}{month:02d}")
        ranges += [[year, month, label, count, url] for label, count in data["priceRanges"]]
        print(f"{year}-{month:02d}", flush=True)
    changed_detached = replace_rows(workbook["Detached_Stats"], "DetachedStatsTable", detached)
    changed_ranges = replace_rows(workbook["Price_Range_Detached"], "PriceRangeDetachedTable", ranges)
    workbook.save(path)
    charts_after = sum(len(ws._charts) for ws in openpyxl.load_workbook(path).worksheets)
    shutil.copyfile(path, PROJECT / "TRREB_Detached_Dataset.xlsx")
    print(f"Detached_Stats: {len(detached)} rows, {changed_detached} changed. "
          f"Price_Range_Detached: {len(ranges)} rows, {changed_ranges} changed. "
          f"Charts before/after: {charts_before}/{charts_after}.")


if __name__ == "__main__":
    main()
